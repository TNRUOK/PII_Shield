"""
extract/text.py
---------------
Universal extractor for plain text, source code, structured data (JSON, CSV, TSV, XML, HTML, YAML),
configuration files (ENV, INI, TOML), log files, markdown, and spreadsheets (XLSX).

Supports any text-decodable file regardless of file extension.
Chunks are partitioned into manageable blocks to optimize NER model processing.
"""
from __future__ import annotations

import csv
import io
import logging
from pathlib import Path
from typing import List, Optional

from extract.models import Chunk

log = logging.getLogger(__name__)

DEFAULT_CHUNK_CHARS = 2500


def is_binary_string(bytes_data: bytes) -> bool:
    """Check if byte sequence appears to be non-text binary data."""
    if not bytes_data:
        return False
    sample = bytes_data[:8192]
    if b"\x00" in sample:
        if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
            return False
        return True
    text_chars = bytearray({7, 8, 9, 10, 12, 13, 27} | set(range(0x20, 0x100)) - {0x7F})
    non_text = sample.translate(None, text_chars)
    return len(non_text) / len(sample) > 0.30


def read_text_file(path: Path) -> Optional[str]:
    """Read a file attempting common encodings."""
    try:
        raw_bytes = path.read_bytes()
    except Exception as exc:
        log.warning("Cannot read bytes from %s: %s", path, exc)
        return None

    for enc in ("utf-8-sig", "utf-8", "utf-16", "latin-1", "cp1252"):
        try:
            return raw_bytes.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return None


def extract_plain_text(file_path: str, config: Optional[dict] = None) -> List[Chunk]:
    """Extract plain text / code / markup into logical chunks."""
    path = Path(file_path)
    text = read_text_file(path)
    if text is None:
        return [
            Chunk(
                file=file_path,
                location="file",
                text="",
                method="blocked",
                ocr_confidence=None,
                order_index=0,
                blocked_reason="Cannot decode file as text",
            )
        ]

    if not text.strip():
        return [
            Chunk(
                file=file_path,
                location="lines:1",
                text="",
                method="native",
                ocr_confidence=None,
                order_index=0,
            )
        ]

    lines = text.splitlines(keepends=True)
    chunks: List[Chunk] = []
    current_lines: List[str] = []
    current_chars = 0
    start_line = 1

    for i, line in enumerate(lines, 1):
        current_lines.append(line)
        current_chars += len(line)

        if current_chars >= DEFAULT_CHUNK_CHARS or (
            current_chars >= 800 and line.strip() == ""
        ):
            chunk_text = "".join(current_lines).strip()
            if chunk_text:
                chunks.append(
                    Chunk(
                        file=file_path,
                        location=f"lines:{start_line}-{i}",
                        text=chunk_text,
                        method="native",
                        ocr_confidence=None,
                        order_index=len(chunks),
                    )
                )
            current_lines = []
            current_chars = 0
            start_line = i + 1

    if current_lines:
        chunk_text = "".join(current_lines).strip()
        if chunk_text:
            chunks.append(
                Chunk(
                    file=file_path,
                    location=f"lines:{start_line}-{len(lines)}",
                    text=chunk_text,
                    method="native",
                    ocr_confidence=None,
                    order_index=len(chunks),
                )
            )

    return chunks if chunks else [
        Chunk(
            file=file_path,
            location="lines:1",
            text=text,
            method="native",
            ocr_confidence=None,
            order_index=0,
        )
    ]


def extract_spreadsheet(file_path: str, config: Optional[dict] = None) -> List[Chunk]:
    """Extract text from XLSX spreadsheet files using openpyxl."""
    path = Path(file_path)
    chunks: List[Chunk] = []

    try:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            sheet_rows = []
            row_idx = 1
            for row in sheet.iter_rows(values_only=True):
                non_empty = [str(val).strip() for val in row if val is not None and str(val).strip()]
                if non_empty:
                    sheet_rows.append(f"Row {row_idx}: " + " | ".join(non_empty))
                row_idx += 1

            if sheet_rows:
                batch_size = 40
                for b_idx in range(0, len(sheet_rows), batch_size):
                    batch = sheet_rows[b_idx : b_idx + batch_size]
                    chunks.append(
                        Chunk(
                            file=file_path,
                            location=f"sheet:{sheet_name}/rows:{b_idx+1}-{b_idx+len(batch)}",
                            text="\n".join(batch),
                            method="native",
                            ocr_confidence=None,
                            order_index=len(chunks),
                        )
                    )
        wb.close()
    except Exception as exc:
        log.warning("openpyxl failed for %s (%s); falling back to raw text", path.name, exc)
        return extract_plain_text(file_path, config)

    if not chunks:
        chunks.append(
            Chunk(
                file=file_path,
                location="sheet:1",
                text="",
                method="native",
                ocr_confidence=None,
                order_index=0,
            )
        )
    return chunks
